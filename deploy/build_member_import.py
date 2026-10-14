#!/usr/bin/env python3
"""把 BLIA 匯出的會員名單轉成投票系統可直接匯入的格式。

來源：分會幹部選舉用會員名單 YYYY-MM-DD.xlsx
      （工作表 All Members = 原始完整名單；Sheet1 = 已篩選的 Active 名單）

輸出欄位（對應 backend/app/routers/members.py 的 HEADER_ALIASES）：
    佛光會員卡號 / 姓名 / givenname / surname / 所屬分會 /
    手機號 / email / 地址 / 是否繳納2026年會費

會費規則（使用者定義）：
    Fee Paid Until 為「會費繳至該年 12/31」。
    >= 2026-12-31  → 2026 年會費已繳 → 「是」
    <  2026-12-31  → 未繳（或只有 2025 以前）→ 「否」

用法：
    python3 deploy/build_member_import.py <來源.xlsx> [輸出.xlsx] [--sheet Sheet1]
"""
from __future__ import annotations

import argparse
import datetime
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

# 匯入檔的表頭（必須是 HEADER_ALIASES 認得的名稱）
OUT_HEADERS = [
    "佛光會員卡號",
    "姓名",
    "givenname",
    "surname",
    "所屬分會",
    "手機號",
    "email",
    "地址",
    "是否繳納2026年會費",
]

# 來源 Subchapter → 系統分會名稱
# （系統以 code 為準：east/south/west/north/central；菩提分會 = north）
SUBCHAPTER_TO_DIVISION = [
    ("青年團", "渥太華佛光青年團"),   # BLIA YAD；須先於「中區」等比對
    ("YAD", "渥太華佛光青年團"),
    ("東區", "東區分會"),
    ("南區", "南區分會"),
    ("西區", "西區分會"),
    ("菩提", "菩提分會"),
    ("中區", "中區分會"),
]

DUES_CUTOFF = datetime.date(2026, 12, 31)


def cell(v: object) -> str:
    """把儲存格轉成乾淨字串（去掉尾端空白與換行）。"""
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date().isoformat()
    if isinstance(v, datetime.date):
        return v.isoformat()
    return str(v).replace("\u3000", " ").strip().strip("\n").strip()


def map_division(subchapter: str) -> str | None:
    """來源 Subchapter → 系統分會名稱；空白或無法對應回 None。"""
    s = subchapter or ""
    if not s:
        return None
    for key, name in SUBCHAPTER_TO_DIVISION:
        if key in s:
            return name
    return None


def dues_paid_2026(value: object) -> str:
    d = value.date() if isinstance(value, datetime.datetime) else value
    if not isinstance(d, datetime.date):
        return "否"
    return "是" if d >= DUES_CUTOFF else "否"


def build(src: Path, dst: Path, sheet: str) -> int:
    wb = openpyxl.load_workbook(src, data_only=True)
    if sheet not in wb.sheetnames:
        raise SystemExit(f"找不到工作表 {sheet!r}；可用的有 {wb.sheetnames}")
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[0]

    def col(key: str) -> int | None:
        for i, h in enumerate(header):
            if h and key in str(h):
                return i
        return None

    C = {
        "no": col("Original Member Card"),
        "name": col("Chinese Name"),
        "given": col("Given Name"),
        "surname": col("Surname"),
        "sub": col("Subchapter"),
        "addr": col("Address 1"),
        "email": col("Email"),
        "home": col("Home Phone"),
        "mobile": col("Mobile Phone"),
        "fees": col("Fees Paid"),
    }
    missing = [k for k, v in C.items() if v is None]
    if missing:
        raise SystemExit(f"來源缺少欄位：{missing}")

    out_rows: list[list[str]] = []
    dropped: list[tuple[str, str]] = []   # (卡號或列號, 原因)
    seen: set[str] = set()

    for n, r in enumerate(rows[1:], start=2):
        no = cell(r[C["no"]])
        name = cell(r[C["name"]])
        sub = cell(r[C["sub"]])
        if not any(cell(r[i]) for i in C.values()):
            continue                                    # 整列空白
        if not no:
            dropped.append((f"第 {n} 列", "沒有會員卡號")); continue
        if not name:
            dropped.append((no, "沒有中文姓名")); continue
        if no in seen:
            dropped.append((no, "卡號重複")); continue
        div = map_division(sub)
        if div is None:
            reason = f"分會空白或無法對應（{sub or '空白'}）"
            dropped.append((no, reason)); continue

        seen.add(no)
        phone = cell(r[C["mobile"]]) or cell(r[C["home"]])
        out_rows.append([
            no, name,
            cell(r[C["given"]]), cell(r[C["surname"]]),
            div, phone, cell(r[C["email"]]), cell(r[C["addr"]]),
            dues_paid_2026(r[C["fees"]]),
        ])

    # ── 寫出 ──
    out = openpyxl.Workbook()
    sh = out.active
    sh.title = "members"
    sh.append(OUT_HEADERS)
    for r in out_rows:
        sh.append(r)
    bold = Font(bold=True)
    fill = PatternFill("solid", fgColor="F2E8D5")
    for c in sh[1]:
        c.font = bold
        c.fill = fill
    widths = [24, 12, 12, 12, 14, 16, 30, 34, 18]
    for i, w in enumerate(widths, start=1):
        sh.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    sh.freeze_panes = "A2"
    dst.parent.mkdir(parents=True, exist_ok=True)
    out.save(dst)

    # ── 報告 ──
    from collections import Counter
    print(f"來源：{src.name}  工作表「{sheet}」")
    print(f"輸出：{dst}")
    print(f"  可匯入 {len(out_rows)} 筆")
    print(f"  已排除 {len(dropped)} 筆")
    if dropped:
        for reason, cnt in Counter(x[1] for x in dropped).most_common():
            print(f"    · {reason}：{cnt}")
    print("  分會分布：")
    for k, v in Counter(r[4] for r in out_rows).most_common():
        print(f"    {v:4d}  {k}")
    paid = sum(1 for r in out_rows if r[8] == "是")
    print(f"  會費：已繳 2026 {paid}／未繳 {len(out_rows) - paid}")
    return len(out_rows)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("src", type=Path)
    ap.add_argument("dst", type=Path, nargs="?", default=Path("deploy/members-import.xlsx"))
    ap.add_argument("--sheet", default="Sheet1",
                    help="工作表名稱（預設 Sheet1＝已篩選的 Active 名單；用 'All Members' 取完整名單）")
    a = ap.parse_args()
    if not a.src.exists():
        raise SystemExit(f"找不到來源檔案：{a.src}")
    build(a.src, a.dst, a.sheet)


if __name__ == "__main__":
    main()
