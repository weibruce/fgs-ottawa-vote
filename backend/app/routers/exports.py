"""資料匯出路由（需管理員認證）

契約：docs/05_api_contract.md 第 8 節
- GET  /admin/exports/history          匯出歷史（最近 50 筆）
- POST /admin/exports/{kind}           產生檔案 + 寫入 export_logs
- GET  /admin/exports/download/{log_id} 以原參數重新產生檔案

CSV 一律加 UTF-8 BOM（Excel 開中文才正常）；檔名放 Content-Disposition，
中文用 filename*=UTF-8''<urlencoded>（前端 exports.ts 依此解析）。
匿名規則：請求帶 anonymous=true 且該輪次 anonymous=true → 不含姓名/卡號/是否代投。
"""
import csv
import io
import json
import urllib.parse
from datetime import datetime

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy import func
from sqlalchemy.orm import Session

from app.database import get_db
from app.deps import get_current_admin
from app.models import (
    Appointment,
    Candidate,
    Division,
    ExportLog,
    Member,
    Round,
    RoundCandidate,
    Vote,
    VoteCandidate,
)
from app.schemas.export import ExportLogOut

router = APIRouter(prefix="/admin/exports", tags=["admin-exports"])

CSV_MEDIA = "text/csv; charset=utf-8"
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# kind → 允許的格式（契約表格）
KIND_FORMATS: dict[str, tuple[str, ...]] = {
    "division_votes": ("csv", "xlsx"),
    "round2_votes": ("csv", "xlsx"),
    "appointments": ("xlsx",),
    "division_summary": ("xlsx",),
    "members": ("csv", "xlsx"),
    "full_report": ("xlsx",),
}


# ─────────────────────────── 共用小工具 ───────────────────────────

def _fmt_dt(dt: datetime | None) -> str:
    """時間 → 本地時區 'YYYY-MM-DD HH:MM:SS'（空值 → 空字串）"""
    if dt is None:
        return ""
    if dt.tzinfo is not None:
        dt = dt.astimezone()
    return dt.strftime("%Y-%m-%d %H:%M:%S")


def _to_csv(headers: list[str], rows: list[list]) -> bytes:
    """單表 → CSV bytes（含 UTF-8 BOM、CRLF，Excel 相容）"""
    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    writer.writerow(headers)
    writer.writerows(rows)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")


def _to_xlsx(sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    """多表 → xlsx bytes（首列粗體、自動欄寬）"""
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title=title[:31])
        ws.append(list(headers))
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append(["" if v is None else v for v in row])
        # 欄寬（依前 200 列估）
        for idx, head in enumerate(headers, start=1):
            sample = [len(str(head))] + [len(str(r[idx - 1])) for r in rows[:200] if len(r) >= idx]
            ws.column_dimensions[get_column_letter(idx)].width = min(max(max(sample) + 2, 8), 40)
    bio = io.BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _content_disposition(filename: str, ext: str) -> str:
    """attachment + RFC 5987 中文檔名（filename*=UTF-8''...）"""
    quoted = urllib.parse.quote(filename, safe="")
    fallback = filename.encode("ascii", "ignore").decode().strip()
    if fallback in ("", "." + ext):
        fallback = f"export.{ext}"
    return f"attachment; filename=\"{fallback}\"; filename*=UTF-8''{quoted}"


def _stream(content: bytes, filename: str, ext: str) -> StreamingResponse:
    media = XLSX_MEDIA if ext == "xlsx" else CSV_MEDIA
    return StreamingResponse(
        io.BytesIO(content),
        media_type=media,
        headers={
            "Content-Disposition": _content_disposition(filename, ext),
            "Content-Length": str(len(content)),
        },
    )


# ─────────────────────────── 資料解析 ───────────────────────────

def _get_division(db: Session, division_id: int | None) -> Division | None:
    """驗證並取得分區（未給 → None；不存在 → 404）"""
    if division_id is None:
        return None
    div = db.get(Division, division_id)
    if div is None:
        raise HTTPException(status_code=404, detail="分區不存在")
    return div


def _resolve_round(db: Session, kind: str, round_id: int | None) -> Round | None:
    """決定匯出使用的輪次。

    - 有帶 round_id → 必須存在
    - round2_votes → 優先 round_no=2
    - 其餘 → active 輪次，其次最新輪次（皆無 → None，輸出空表）
    """
    if round_id is not None:
        rnd = db.get(Round, round_id)
        if rnd is None:
            raise HTTPException(status_code=404, detail="輪次不存在")
        return rnd
    if kind == "round2_votes":
        rnd = db.query(Round).filter(Round.round_no == 2).order_by(Round.id.desc()).first()
        if rnd is not None:
            return rnd
    return (
        db.query(Round).filter(Round.status == "active").order_by(Round.id.desc()).first()
        or db.query(Round).order_by(Round.id.desc()).first()
    )


def _division_rows(db: Session, division_id: int | None) -> list[Division]:
    q = db.query(Division).filter(Division.is_active == True)  # noqa: E712
    if division_id is not None:
        q = q.filter(Division.id == division_id)
    return q.order_by(Division.sort_order, Division.id).all()


def _vote_candidate_names(db: Session, vote_ids: list[int]) -> dict[int, list[str]]:
    """vote_id → 投了誰（候選人名，依 sort_order）"""
    if not vote_ids:
        return {}
    rows = (
        db.query(VoteCandidate.vote_id, Candidate.name)
        .join(Candidate, Candidate.id == VoteCandidate.candidate_id)
        .filter(VoteCandidate.vote_id.in_(vote_ids))
        .order_by(Candidate.sort_order, Candidate.id)
        .all()
    )
    out: dict[int, list[str]] = {}
    for vote_id, name in rows:
        out.setdefault(vote_id, []).append(name)
    return out


def _vote_detail(
    db: Session,
    rnd: Round | None,
    division_id: int | None,
    req_anonymous: bool,
    include_round: bool = False,
) -> tuple[list[str], list[list]]:
    """投票明細：姓名/卡號/所屬分區/是否代投/投票時間/投了誰

    匿名判定：`req_anonymous` 且該輪次 `anonymous` 皆為真 → 去掉姓名/卡號/是否代投。
    """
    masked = bool(req_anonymous and rnd is not None and rnd.anonymous)
    headers = ["所屬分區", "投票時間", "投了誰"] if masked else [
        "姓名", "卡號", "所屬分區", "是否代投", "投票時間", "投了誰",
    ]
    if include_round:
        headers = ["輪次"] + headers

    if rnd is None:
        return headers, []
    q = db.query(Vote).filter(Vote.round_id == rnd.id)
    if division_id is not None:
        q = q.filter(Vote.division_id == division_id)
    votes = q.order_by(Vote.created_at, Vote.id).all()

    names = _vote_candidate_names(db, [v.id for v in votes])
    div_names = {d.id: d.name for d in db.query(Division).all()}

    rows: list[list] = []
    for v in votes:
        who = "、".join(names.get(v.id, []))
        if masked:
            row: list = [div_names.get(v.division_id, ""), _fmt_dt(v.created_at), who]
        else:
            row = [
                v.member_name,
                v.member_no,
                div_names.get(v.division_id, ""),
                "是" if v.is_proxy else "否",
                _fmt_dt(v.created_at),
                who,
            ]
        if include_round:
            row = [rnd.name] + row
        rows.append(row)
    return headers, rows


def _candidate_vote_counts(db: Session, rnd: Round | None, division_id: int) -> dict[str, int]:
    """某輪次某分區：候選人名 → 得票數"""
    if rnd is None:
        return {}
    rows = (
        db.query(Candidate.name, func.count(VoteCandidate.id))
        .join(VoteCandidate, VoteCandidate.candidate_id == Candidate.id)
        .join(Vote, Vote.id == VoteCandidate.vote_id)
        .filter(Vote.round_id == rnd.id, Candidate.division_id == division_id)
        .group_by(Candidate.id, Candidate.name)
        .all()
    )
    return {name: count for name, count in rows}


def _division_summary(
    db: Session, rnd: Round | None, division_id: int | None
) -> tuple[list[str], list[list]]:
    """五區彙總：會員數/已投票/投票率/候選人數/代投票數/最高票/平票"""
    headers = [
        "分區", "會員數", "已投票", "投票率", "候選人數",
        "代投票數", "最高票候選人", "最高票數", "是否平票",
    ]
    rows: list[list] = []
    for div in _division_rows(db, division_id):
        total_members = db.query(func.count(Member.id)).filter(
            Member.division_id == div.id, Member.is_active == True  # noqa: E712
        ).scalar() or 0

        if rnd is not None:
            votes = db.query(Vote).filter(
                Vote.division_id == div.id, Vote.round_id == rnd.id
            ).all()
            cand_count = db.query(func.count(RoundCandidate.id)).filter(
                RoundCandidate.round_id == rnd.id, RoundCandidate.division_id == div.id
            ).scalar() or 0
        else:
            votes = []
            cand_count = db.query(func.count(Candidate.id)).filter(
                Candidate.division_id == div.id, Candidate.is_active == True  # noqa: E712
            ).scalar() or 0

        voted = len(votes)
        proxy = sum(1 for v in votes if v.is_proxy)
        rate = round(voted / total_members * 100, 1) if total_members else 0.0

        counts = _candidate_vote_counts(db, rnd, div.id)
        if counts:
            top_votes = max(counts.values())
            tops = sorted([n for n, c in counts.items() if c == top_votes])
            top_name = "、".join(tops)
            is_tie = len(tops) > 1 and top_votes > 0
        else:
            top_votes, top_name, is_tie = 0, "", False

        rows.append([
            div.name, total_members, voted, f"{rate}%", cand_count, proxy,
            top_name, top_votes, "是" if is_tie else "否",
        ])
    return headers, rows


def _appointments_rows(
    db: Session, division_id: int | None
) -> tuple[list[str], list[list]]:
    """幹部指派名單"""
    headers = ["分區", "職務", "姓名", "任期", "任命人", "是否確認", "建立時間"]
    q = db.query(Appointment)
    if division_id is not None:
        q = q.filter(Appointment.division_id == division_id)
    appts = q.order_by(Appointment.division_id, Appointment.id).all()
    div_names = {d.id: d.name for d in db.query(Division).all()}
    rows = [
        [
            div_names.get(a.division_id, ""),
            a.position,
            a.name,
            a.term,
            a.appointed_by,
            "是" if a.is_confirmed else "否",
            _fmt_dt(a.created_at),
        ]
        for a in appts
    ]
    return headers, rows


def _members_rows(
    db: Session, rnd: Round | None, division_id: int | None
) -> tuple[list[str], list[list]]:
    """會員名單（含該輪次是否已投票）"""
    headers = [
        "佛光會員卡號", "姓名(繁)", "姓名(簡)", "givenname", "surname", "所屬分會",
        "性別", "手機號", "Email", "地址", "已投票", "投票時間", "狀態",
    ]
    q = db.query(Member)
    if division_id is not None:
        q = q.filter(Member.division_id == division_id)
    members = q.order_by(Member.division_id, Member.member_no).all()

    voted_at: dict[str, datetime | None] = {}
    if rnd is not None:
        for vote in db.query(Vote).filter(Vote.round_id == rnd.id).all():
            voted_at.setdefault(vote.member_no, vote.created_at)

    div_names = {d.id: d.name for d in db.query(Division).all()}
    rows = []
    for m in members:
        at = voted_at.get(m.member_no)
        rows.append([
            m.member_no,
            m.name_trad,
            m.name_simp,
            m.givenname,
            m.surname,
            div_names.get(m.division_id, ""),
            m.gender,
            m.phone,
            m.email,
            m.address,
            "是" if at is not None else "否",
            _fmt_dt(at) if at is not None else "",
            "啟用" if m.is_active else "停用",
        ])
    return headers, rows


# ─────────────────────────── 檔案組裝 ───────────────────────────

def _collect(
    db: Session, kind: str, division_id: int | None, round_id: int | None, anonymous: bool
) -> tuple[str, list[tuple[str, list[str], list[list]]]]:
    """回傳 (檔名（不含副檔名）, sheets)"""
    div = _get_division(db, division_id)
    prefix = div.name if div is not None else ""

    if kind in ("division_votes", "round2_votes"):
        rnd = _resolve_round(db, kind, round_id)
        headers, rows = _vote_detail(db, rnd, division_id, anonymous)
        if kind == "round2_votes":
            stem = f"{prefix}第二輪投票明細" if prefix else "第二輪投票明細"
        else:
            stem = f"{prefix}投票明細" if prefix else "各分區投票明細"
        return stem, [("投票明細", headers, rows)]

    if kind == "division_summary":
        rnd = _resolve_round(db, kind, round_id)
        headers, rows = _division_summary(db, rnd, division_id)
        return "五區彙總統計", [("彙總統計", headers, rows)]

    if kind == "appointments":
        headers, rows = _appointments_rows(db, division_id)
        stem = f"{prefix}幹部指派名單" if prefix else "幹部指派名單"
        return stem, [("幹部指派", headers, rows)]

    if kind == "members":
        rnd = _resolve_round(db, kind, round_id)
        headers, rows = _members_rows(db, rnd, division_id)
        stem = f"{prefix}會員名單" if prefix else "會員名單"
        return stem, [("會員名單", headers, rows)]

    if kind == "full_report":
        rounds = db.query(Round).order_by(Round.round_no, Round.id).all()

        summary_headers = [
            "輪次", "分區", "會員數", "已投票", "投票率", "候選人數",
            "代投票數", "最高票候選人", "最高票數", "是否平票",
        ]
        summary_rows: list[list] = []
        vote_headers = ["輪次"]
        vote_rows: list[list] = []
        for rnd in rounds:
            _, rows = _division_summary(db, rnd, division_id)
            summary_rows.extend([[rnd.name] + r for r in rows])
            headers, rows = _vote_detail(db, rnd, division_id, anonymous, include_round=True)
            vote_headers = headers
            vote_rows.extend(rows)

        appt_headers, appt_rows = _appointments_rows(db, division_id)
        member_rnd = _resolve_round(db, "members", round_id)
        member_headers, member_rows = _members_rows(db, member_rnd, division_id)

        sheets = [
            ("彙總統計", summary_headers, summary_rows),
            ("投票明細", vote_headers, vote_rows),
            ("幹部指派", appt_headers, appt_rows),
            ("會員名單", member_headers, member_rows),
        ]
        return "完整選舉報告", sheets

    raise HTTPException(status_code=404, detail=f"不支援的匯出類型：{kind}")


def _generate(
    db: Session, kind: str, params: dict
) -> tuple[bytes, str, str]:
    """依參數產生檔案 → (content, filename, ext)"""
    ext = params.get("format", "csv")
    if ext not in KIND_FORMATS[kind]:
        raise HTTPException(status_code=400, detail=f"{kind} 不支援 {ext} 格式")
    stem, sheets = _collect(
        db,
        kind,
        params.get("division_id"),
        params.get("round_id"),
        bool(params.get("anonymous")),
    )
    filename = f"{stem}.{ext}"
    if ext == "xlsx":
        content = _to_xlsx(sheets)
    else:
        content = _to_csv(sheets[0][1], sheets[0][2])
    return content, filename, ext


# ─────────────────────────── 端點 ───────────────────────────

@router.get("/history", response_model=list[ExportLogOut])
def export_history(db: Session = Depends(get_db), _admin=Depends(get_current_admin)):
    """最近 50 筆匯出紀錄（新 → 舊）"""
    logs = (
        db.query(ExportLog)
        .order_by(ExportLog.created_at.desc(), ExportLog.id.desc())
        .limit(50)
        .all()
    )
    return [
        ExportLogOut(
            id=log.id,
            kind=log.kind,
            filename=log.filename,
            size=log.size,
            format=log.file_format,
            operator=log.operator,
            created_at=log.created_at,
        )
        for log in logs
    ]


@router.get("/download/{log_id}")
def download_export(
    log_id: int,
    db: Session = Depends(get_db),
    _admin=Depends(get_current_admin),
):
    """以 export_logs.params 的原參數重新產生同一份檔案"""
    log = db.get(ExportLog, log_id)
    if log is None:
        raise HTTPException(status_code=404, detail="匯出紀錄不存在")
    if log.kind not in KIND_FORMATS:
        raise HTTPException(status_code=400, detail=f"不支援的匯出類型：{log.kind}")

    try:
        params = json.loads(log.params) if log.params else {}
    except (TypeError, ValueError):
        params = {}
    params.setdefault("format", log.file_format or "csv")

    content, _, ext = _generate(db, log.kind, params)
    return _stream(content, log.filename or f"{log.kind}.{ext}", ext)


@router.post("/{kind}")
def create_export(
    kind: str,
    division_id: int | None = Query(None, description="分區 ID（不帶=全部分區）"),
    round_id: int | None = Query(None, description="輪次 ID（不帶=當前/預設輪次）"),
    anonymous: bool = Query(False, description="匿名模式（與輪次 anonymous 同時為真才遮蔽）"),
    format: str = Query("csv", description="csv | xlsx"),
    db: Session = Depends(get_db),
    admin=Depends(get_current_admin),
):
    """產生檔案、寫入 export_logs，並以 StreamingResponse 回傳"""
    if kind not in KIND_FORMATS:
        raise HTTPException(status_code=404, detail=f"不支援的匯出類型：{kind}")
    if format not in ("csv", "xlsx"):
        raise HTTPException(status_code=400, detail="format 僅支援 csv 或 xlsx")

    params = {
        "division_id": division_id,
        "round_id": round_id,
        "anonymous": anonymous,
        "format": format,
    }
    content, filename, ext = _generate(db, kind, params)

    operator = getattr(admin, "display_name", None) or getattr(admin, "username", "admin")
    log = ExportLog(
        kind=kind,
        filename=filename,
        file_format=ext,
        size=len(content),
        params=json.dumps(params, ensure_ascii=False),
        operator=operator,
    )
    db.add(log)
    db.commit()
    db.refresh(log)

    return _stream(content, filename, ext)
